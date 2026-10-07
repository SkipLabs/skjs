"""Build build/images.pptx (each rendered slide full-bleed) and build/notes.txt.

Run after render.py; keynote.applescript then adds the notes in Keynote and
writes out/slides.key and out/slides.pptx. Keynote refuses to import notes
written by python-pptx, so the notes go in through Keynote itself.
"""

import html
import json
import re

from pptx import Presentation
from pptx.util import Emu, Inches

deck = json.load(open("deck/deck.json"))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(40 / 3), Inches(7.5)
blank = prs.slide_layouts[6]
notes = []
for n, sid in enumerate(deck["order"], 1):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(f"out/slides/{n:02d}-{sid}.png", Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    aside = re.search(r"<aside>(.*?)</aside>", open(f"deck/slides/{sid}.html").read(), flags=re.S)
    notes.append(html.unescape(aside.group(1).strip()) if aside else "")
prs.save("build/images.pptx")
open("build/notes.txt", "w").write("\n".join(notes) + "\n")
print(f"build/images.pptx: {len(prs.slides)} slides")
